import os
import calendar
from datetime import date, datetime
from functools import wraps
from io import BytesIO

from flask import (Flask, render_template, request, redirect, url_for,
                   flash, jsonify, send_file, abort)
from flask_login import (LoginManager, login_user, logout_user,
                         login_required, current_user)
from sqlalchemy import extract

from config import Config
from models import db, User, Student, Subject, Department, AttendanceRecord, Condonation


def create_app(config_class=Config):
    app = Flask(__name__)
    app.config.from_object(config_class)

    db.init_app(app)

    login_manager = LoginManager(app)
    login_manager.login_view = 'login'
    login_manager.login_message = 'Please log in to access this page.'
    login_manager.login_message_category = 'info'

    @login_manager.user_loader
    def load_user(user_id):
        return db.session.get(User, int(user_id))

    # ── Role decorators ──────────────────────────────────────────────────────
    def admin_required(f):
        @wraps(f)
        def decorated(*args, **kwargs):
            if not current_user.is_authenticated or current_user.role != 'admin':
                abort(403)
            return f(*args, **kwargs)
        return decorated

    def faculty_or_above(f):
        @wraps(f)
        def decorated(*args, **kwargs):
            if not current_user.is_authenticated or current_user.role not in ('faculty', 'hod', 'admin'):
                abort(403)
            return f(*args, **kwargs)
        return decorated

    def hod_or_above(f):
        @wraps(f)
        def decorated(*args, **kwargs):
            if not current_user.is_authenticated or current_user.role not in ('hod', 'admin'):
                abort(403)
            return f(*args, **kwargs)
        return decorated

    # ── Helpers ─────────────────────────────────────────────────────────────
    def get_monthly_summary(subject_id, month, year):
        subject = db.session.get(Subject, subject_id)
        if not subject:
            return []
        students = (Student.query
                    .filter_by(department_id=subject.department_id,
                               semester=subject.semester,
                               section=subject.section)
                    .order_by(Student.usn).all())
        results = []
        for student in students:
            records = (AttendanceRecord.query
                       .filter(AttendanceRecord.student_id == student.id,
                               AttendanceRecord.subject_id == subject_id,
                               extract('month', AttendanceRecord.date) == month,
                               extract('year',  AttendanceRecord.date) == year)
                       .all())
            total    = len(records)
            attended = sum(1 for r in records if r.status in ('P', 'L'))
            pct      = round((attended / total * 100) if total else 0, 2)

            cond = Condonation.query.filter_by(
                student_id=student.id, subject_id=subject_id,
                month=month, year=year).first()
            condoned = cond.condoned_classes if cond else 0
            eff_att  = attended + condoned
            eff_pct  = round((eff_att / total * 100) if total else 0, 2)

            results.append({
                'student':             student,
                'total':               total,
                'attended':            attended,
                'percentage':          pct,
                'condoned':            condoned,
                'effective_percentage': eff_pct,
                'shortage':            (eff_pct < Config.ATTENDANCE_THRESHOLD and total > 0),
            })
        return results

    app.jinja_env.globals['get_monthly_summary'] = get_monthly_summary
    app.jinja_env.globals['now'] = datetime.now

    # ── Auth ─────────────────────────────────────────────────────────────────
    @app.route('/')
    def index():
        return redirect(url_for('dashboard') if current_user.is_authenticated else url_for('login'))

    @app.route('/login', methods=['GET', 'POST'])
    def login():
        if current_user.is_authenticated:
            return redirect(url_for('dashboard'))
        if request.method == 'POST':
            username = request.form.get('username', '').strip()
            password = request.form.get('password', '')
            user = User.query.filter_by(username=username).first()
            if user and user.check_password(password) and user.is_active_user:
                login_user(user, remember=bool(request.form.get('remember')))
                flash(f'Welcome back, {user.name}!', 'success')
                nxt = request.args.get('next')
                return redirect(nxt or url_for('dashboard'))
            flash('Invalid username or password.', 'danger')
        return render_template('login.html')

    @app.route('/logout')
    @login_required
    def logout():
        logout_user()
        flash('You have been logged out.', 'info')
        return redirect(url_for('login'))

    # ── Dashboard ────────────────────────────────────────────────────────────
    @app.route('/dashboard')
    @login_required
    def dashboard():
        today = date.today()
        month, year = today.month, today.year

        if current_user.role == 'admin':
            departments   = Department.query.all()
            total_students = Student.query.count()
            total_faculty  = User.query.filter_by(role='faculty').count()
            total_subjects = Subject.query.count()
            shortage_count = 0
            for subj in Subject.query.all():
                shortage_count += sum(1 for s in get_monthly_summary(subj.id, month, year) if s['shortage'])
            return render_template('dashboard/admin.html',
                departments=departments, total_students=total_students,
                total_faculty=total_faculty, total_subjects=total_subjects,
                shortage_count=shortage_count, today=today, month=month, year=year)

        elif current_user.role == 'hod':
            dept = current_user.department
            subjects = Subject.query.filter_by(department_id=dept.id).all()
            shortage_ids = set()
            for subj in subjects:
                for s in get_monthly_summary(subj.id, month, year):
                    if s['shortage']:
                        shortage_ids.add(s['student'].id)
            return render_template('dashboard/hod.html',
                department=dept,
                student_count=Student.query.filter_by(department_id=dept.id).count(),
                subject_count=len(subjects),
                shortage_count=len(shortage_ids),
                today=today, month=month, year=year)

        elif current_user.role == 'faculty':
            subjects = Subject.query.filter_by(faculty_id=current_user.id).all()
            today_records = {}
            for subj in subjects:
                total_s = Student.query.filter_by(
                    department_id=subj.department_id,
                    semester=subj.semester, section=subj.section).count()
                marked = AttendanceRecord.query.filter_by(
                    subject_id=subj.id, date=today).count()
                today_records[subj.id] = {
                    'marked': marked, 'total': total_s,
                    'done': (marked >= total_s and total_s > 0)}
            return render_template('dashboard/faculty.html',
                subjects=subjects, today_records=today_records, today=today)

        else:  # student
            student = Student.query.filter_by(user_id=current_user.id).first()
            if not student:
                flash('Student profile not found. Contact admin.', 'danger')
                return redirect(url_for('login'))
            subjects = Subject.query.filter_by(
                department_id=student.department_id,
                semester=student.semester, section=student.section).all()
            att_data = []
            for subj in subjects:
                records = (AttendanceRecord.query
                           .filter(AttendanceRecord.student_id == student.id,
                                   AttendanceRecord.subject_id == subj.id,
                                   extract('month', AttendanceRecord.date) == month,
                                   extract('year',  AttendanceRecord.date) == year)
                           .all())
                total    = len(records)
                attended = sum(1 for r in records if r.status in ('P', 'L'))
                pct      = round((attended / total * 100) if total else 0, 2)
                cond = Condonation.query.filter_by(
                    student_id=student.id, subject_id=subj.id,
                    month=month, year=year).first()
                condoned = cond.condoned_classes if cond else 0
                eff_pct  = round(((attended + condoned) / total * 100) if total else 0, 2)
                att_data.append({
                    'subject': subj, 'total': total, 'attended': attended,
                    'percentage': pct, 'condoned': condoned,
                    'effective_percentage': eff_pct,
                    'shortage': eff_pct < Config.ATTENDANCE_THRESHOLD and total > 0})
            return render_template('dashboard/student.html',
                student=student, att_data=att_data,
                month=month, year=year, month_name=calendar.month_name[month])

    # ── Mark Attendance ──────────────────────────────────────────────────────
    @app.route('/attendance/mark', methods=['GET', 'POST'])
    @login_required
    @faculty_or_above
    def mark_attendance():
        if current_user.role == 'faculty':
            subjects = Subject.query.filter_by(faculty_id=current_user.id).all()
        elif current_user.role == 'hod':
            subjects = Subject.query.filter_by(department_id=current_user.department_id).all()
        else:
            subjects = Subject.query.order_by(Subject.department_id, Subject.semester).all()

        sel_subj_id  = request.args.get('subject_id', type=int)
        sel_date_str = request.args.get('date', date.today().isoformat())
        try:
            sel_date = date.fromisoformat(sel_date_str)
        except ValueError:
            sel_date = date.today()

        students, existing = [], {}
        if sel_subj_id:
            subj = db.session.get(Subject, sel_subj_id)
            if subj:
                students = (Student.query.filter_by(
                    department_id=subj.department_id,
                    semester=subj.semester, section=subj.section)
                    .order_by(Student.usn).all())
                existing = {r.student_id: r.status
                            for r in AttendanceRecord.query.filter_by(
                                subject_id=sel_subj_id, date=sel_date).all()}

        if request.method == 'POST':
            subj_id  = request.form.get('subject_id', type=int)
            date_str = request.form.get('date')
            try:
                att_date = date.fromisoformat(date_str)
            except (ValueError, TypeError):
                att_date = date.today()
            subj = db.session.get(Subject, subj_id)
            stu_list = (Student.query.filter_by(
                department_id=subj.department_id,
                semester=subj.semester, section=subj.section).all()) if subj else []
            for stu in stu_list:
                status = request.form.get(f'status_{stu.id}', 'A')
                rec = AttendanceRecord.query.filter_by(
                    student_id=stu.id, subject_id=subj_id, date=att_date).first()
                if rec:
                    rec.status = status
                    rec.marked_by_id = current_user.id
                    rec.marked_at = datetime.utcnow()
                else:
                    db.session.add(AttendanceRecord(
                        student_id=stu.id, subject_id=subj_id,
                        date=att_date, status=status, marked_by_id=current_user.id))
            db.session.commit()
            flash(f'Attendance saved for {len(stu_list)} students on {att_date.strftime("%d %b %Y")}.', 'success')
            return redirect(url_for('mark_attendance', subject_id=subj_id, date=att_date.isoformat()))

        return render_template('attendance/mark.html',
            subjects=subjects, sel_subj_id=sel_subj_id,
            sel_date=sel_date, students=students, existing=existing,
            today=date.today())

    # ── Monthly Report ───────────────────────────────────────────────────────
    @app.route('/attendance/monthly')
    @login_required
    def monthly_attendance():
        today = date.today()
        month      = request.args.get('month',    today.month, type=int)
        year       = request.args.get('year',     today.year,  type=int)
        dept_id    = request.args.get('dept_id',  type=int)
        semester   = request.args.get('semester', type=int)
        subject_id = request.args.get('subject_id', type=int)

        if current_user.role == 'faculty':
            subjects = Subject.query.filter_by(faculty_id=current_user.id).all()
        elif current_user.role == 'hod':
            dept_id  = current_user.department_id
            subjects = Subject.query.filter_by(department_id=dept_id).all()
        elif current_user.role == 'student':
            stu = Student.query.filter_by(user_id=current_user.id).first()
            subjects = (Subject.query.filter_by(
                department_id=stu.department_id,
                semester=stu.semester, section=stu.section).all()) if stu else []
        else:
            subjects = (Subject.query.filter_by(department_id=dept_id).all()
                        if dept_id else Subject.query.all())

        if semester:
            subjects = [s for s in subjects if s.semester == semester]

        summary = get_monthly_summary(subject_id, month, year) if subject_id else []

        return render_template('attendance/monthly.html',
            departments=Department.query.order_by(Department.name).all(),
            subjects=subjects, summary=summary,
            month=month, year=year, dept_id=dept_id,
            semester=semester, subject_id=subject_id,
            months=[(i, calendar.month_name[i]) for i in range(1, 13)],
            years=list(range(2022, today.year + 2)),
            month_name=calendar.month_name[month],
            threshold=Config.ATTENDANCE_THRESHOLD)

    # ── Student Detail ───────────────────────────────────────────────────────
    @app.route('/attendance/student/<int:student_id>')
    @login_required
    def student_detail(student_id):
        student = db.session.get(Student, student_id) or abort(404)
        if current_user.role == 'student':
            own = Student.query.filter_by(user_id=current_user.id).first()
            if not own or own.id != student_id:
                abort(403)
        today = date.today()
        month = request.args.get('month', today.month, type=int)
        year  = request.args.get('year',  today.year,  type=int)
        subjects = Subject.query.filter_by(
            department_id=student.department_id,
            semester=student.semester, section=student.section).all()
        att_data = []
        for subj in subjects:
            records = (AttendanceRecord.query
                       .filter(AttendanceRecord.student_id == student.id,
                               AttendanceRecord.subject_id == subj.id,
                               extract('month', AttendanceRecord.date) == month,
                               extract('year',  AttendanceRecord.date) == year)
                       .order_by(AttendanceRecord.date).all())
            total    = len(records)
            attended = sum(1 for r in records if r.status in ('P', 'L'))
            pct      = round((attended / total * 100) if total else 0, 2)
            cond = Condonation.query.filter_by(
                student_id=student.id, subject_id=subj.id, month=month, year=year).first()
            condoned = cond.condoned_classes if cond else 0
            eff_pct  = round(((attended + condoned) / total * 100) if total else 0, 2)
            att_data.append({
                'subject': subj, 'records': records, 'total': total,
                'attended': attended, 'percentage': pct,
                'condoned': condoned, 'effective_percentage': eff_pct,
                'shortage': eff_pct < Config.ATTENDANCE_THRESHOLD and total > 0})
        return render_template('students/detail.html',
            student=student, att_data=att_data,
            month=month, year=year, month_name=calendar.month_name[month],
            months=[(i, calendar.month_name[i]) for i in range(1, 13)],
            years=list(range(2022, today.year + 2)),
            threshold=Config.ATTENDANCE_THRESHOLD)

    # ── Students List ────────────────────────────────────────────────────────
    @app.route('/students')
    @login_required
    def students_list():
        if current_user.role == 'student':
            abort(403)
        dept_id  = request.args.get('dept_id',  type=int)
        semester = request.args.get('semester', type=int)
        section  = request.args.get('section')
        q        = request.args.get('q', '').strip()
        if current_user.role in ('hod', 'faculty'):
            dept_id = current_user.department_id
        qry = Student.query
        if dept_id:  qry = qry.filter_by(department_id=dept_id)
        if semester: qry = qry.filter_by(semester=semester)
        if section:  qry = qry.filter_by(section=section)
        if q:
            qry = qry.filter(
                (Student.name.ilike(f'%{q}%')) | (Student.usn.ilike(f'%{q}%')))
        students = qry.order_by(Student.department_id, Student.semester, Student.usn).all()
        return render_template('students/list.html',
            students=students,
            departments=Department.query.order_by(Department.name).all(),
            dept_id=dept_id, semester=semester, section=section, q=q)

    # ── Condonation ──────────────────────────────────────────────────────────
    @app.route('/condonation', methods=['GET', 'POST'])
    @login_required
    @hod_or_above
    def condonation():
        today = date.today()
        month   = request.args.get('month',   today.month, type=int)
        year    = request.args.get('year',    today.year,  type=int)
        dept_id = (current_user.department_id if current_user.role == 'hod'
                   else request.args.get('dept_id', type=int))

        if request.method == 'POST':
            stu_id   = request.form.get('student_id',      type=int)
            subj_id  = request.form.get('subject_id',      type=int)
            classes  = request.form.get('condoned_classes', 0, type=int)
            reason   = request.form.get('reason', '')
            m        = request.form.get('month',  type=int)
            y        = request.form.get('year',   type=int)
            rec = Condonation.query.filter_by(
                student_id=stu_id, subject_id=subj_id, month=m, year=y).first()
            if rec:
                rec.condoned_classes = classes
                rec.reason = reason
                rec.condoned_by_id = current_user.id
            else:
                db.session.add(Condonation(
                    student_id=stu_id, subject_id=subj_id,
                    month=m, year=y, condoned_classes=classes,
                    reason=reason, condoned_by_id=current_user.id))
            db.session.commit()
            flash('Condonation updated successfully.', 'success')
            return redirect(url_for('condonation', month=m, year=y, dept_id=dept_id))

        shortage_list = []
        if dept_id:
            seen = set()
            for subj in Subject.query.filter_by(department_id=dept_id).all():
                for s in get_monthly_summary(subj.id, month, year):
                    key = (s['student'].id, subj.id)
                    if key not in seen and s['shortage']:
                        seen.add(key)
                        shortage_list.append({**s, 'subject': subj})

        return render_template('attendance/condonation.html',
            departments=Department.query.all(),
            shortage_list=shortage_list,
            month=month, year=year, dept_id=dept_id,
            months=[(i, calendar.month_name[i]) for i in range(1, 13)],
            years=list(range(2022, today.year + 2)),
            month_name=calendar.month_name[month],
            threshold=Config.ATTENDANCE_THRESHOLD)

    # ── Excel Export ─────────────────────────────────────────────────────────
    @app.route('/export/excel/<int:subject_id>/<int:month>/<int:year>')
    @login_required
    def export_excel(subject_id, month, year):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        subj    = db.session.get(Subject, subject_id) or abort(404)
        summary = get_monthly_summary(subject_id, month, year)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f'{calendar.month_abbr[month]} {year}'

        hdr_fill = PatternFill('solid', fgColor='1F4E79')
        hdr_font = Font(bold=True, color='FFFFFF', size=10)
        center   = Alignment(horizontal='center', vertical='center')

        def mcell(row, col, val, bold=False, bg=None, color='000000', sz=10):
            c = ws.cell(row=row, column=col, value=val)
            c.font = Font(bold=bold, color=color, size=sz)
            c.alignment = center
            if bg:
                c.fill = PatternFill('solid', fgColor=bg)
            return c

        ws.merge_cells('A1:I1')
        mcell(1,1, Config.COLLEGE_NAME, bold=True, sz=14)
        ws.merge_cells('A2:I2')
        mcell(2,1, f'Monthly Attendance Report — {calendar.month_name[month]} {year}', bold=True, sz=12)
        ws.merge_cells('A3:I3')
        mcell(3,1, f'Subject: {subj.name} ({subj.code})  |  Dept: {subj.department.name}  |  Sem: {subj.semester}  |  Sec: {subj.section}', sz=10)
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 20

        headers = ['Sl.', 'USN', 'Student Name', 'Total Classes', 'Attended', '%', 'Condoned', 'Eff. %', 'Status']
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=5, column=c, value=h)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = center
        ws.row_dimensions[5].height = 18

        red_fill   = PatternFill('solid', fgColor='FFD6D6')
        green_fill = PatternFill('solid', fgColor='D6FFD6')

        for i, row in enumerate(summary, 1):
            stu = row['student']
            fill = red_fill if row['shortage'] else green_fill
            vals = [i, stu.usn, stu.name, row['total'], row['attended'],
                    f"{row['percentage']}%", row['condoned'],
                    f"{row['effective_percentage']}%",
                    'SHORT' if row['shortage'] else 'OK']
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=5+i, column=c, value=v)
                cell.fill = fill
                cell.alignment = center

        for c, w in zip(range(1,10), [6,18,32,14,12,10,12,10,10]):
            ws.column_dimensions[get_column_letter(c)].width = w

        shortage = sum(1 for s in summary if s['shortage'])
        sr = 7 + len(summary)
        ws.cell(row=sr, column=1, value='Summary').font = Font(bold=True)
        ws.cell(row=sr, column=2, value=f'Total: {len(summary)}')
        ws.cell(row=sr, column=3, value=f'Shortage (<{Config.ATTENDANCE_THRESHOLD}%): {shortage}')

        buf = BytesIO()
        wb.save(buf); buf.seek(0)
        fname = f'Attendance_{subj.code}_{calendar.month_name[month]}_{year}.xlsx'
        return send_file(buf, as_attachment=True, download_name=fname,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    # ── PDF Export ───────────────────────────────────────────────────────────
    @app.route('/export/pdf/<int:subject_id>/<int:month>/<int:year>')
    @login_required
    def export_pdf(subject_id, month, year):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib.enums import TA_CENTER
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

        subj    = db.session.get(Subject, subject_id) or abort(404)
        summary = get_monthly_summary(subject_id, month, year)

        buf = BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                rightMargin=1*cm, leftMargin=1*cm,
                                topMargin=1.5*cm, bottomMargin=1*cm)
        styles  = getSampleStyleSheet()
        t_style = ParagraphStyle('T', parent=styles['Title'], fontSize=15, alignment=TA_CENTER, spaceAfter=4)
        s_style = ParagraphStyle('S', parent=styles['Normal'], fontSize=10, alignment=TA_CENTER, spaceAfter=4)

        elems = [
            Paragraph(Config.COLLEGE_NAME, t_style),
            Paragraph(f'Monthly Attendance Report — {calendar.month_name[month]} {year}', s_style),
            Paragraph(f'Subject: {subj.name} ({subj.code})  |  Dept: {subj.department.name}  |  Sem: {subj.semester}  |  Section: {subj.section}', s_style),
            Spacer(1, 0.4*cm),
        ]

        data = [['Sl.', 'USN', 'Name', 'Total', 'Attended', '%', 'Condoned', 'Eff. %', 'Status']]
        for i, row in enumerate(summary, 1):
            s = row['student']
            data.append([str(i), s.usn, s.name, str(row['total']), str(row['attended']),
                         f"{row['percentage']}%", str(row['condoned']),
                         f"{row['effective_percentage']}%",
                         'SHORT' if row['shortage'] else 'OK'])

        col_w = [1.2*cm, 3.5*cm, 6*cm, 2*cm, 2.5*cm, 2*cm, 2.5*cm, 2*cm, 2*cm]
        tbl   = Table(data, colWidths=col_w, repeatRows=1)
        cmds  = [
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E79')),
            ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
            ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE',   (0,0), (-1,0), 9),
            ('FONTSIZE',   (0,1), (-1,-1), 8),
            ('ALIGN',      (0,0), (-1,-1), 'CENTER'),
            ('VALIGN',     (0,0), (-1,-1), 'MIDDLE'),
            ('GRID',       (0,0), (-1,-1), 0.4, colors.black),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F5F5F5')]),
        ]
        for i, row in enumerate(summary, 1):
            if row['shortage']:
                cmds.append(('BACKGROUND', (0,i), (-1,i), colors.HexColor('#FFD6D6')))
                cmds.append(('TEXTCOLOR',  (8,i), (8,i),  colors.red))
        tbl.setStyle(TableStyle(cmds))
        elems.append(tbl)

        shortage = sum(1 for s in summary if s['shortage'])
        elems.append(Spacer(1, 0.5*cm))
        elems.append(Paragraph(
            f'Total Students: {len(summary)}  |  Shortage (< {Config.ATTENDANCE_THRESHOLD}%): {shortage}  |  Generated: {datetime.now().strftime("%d %b %Y %H:%M")}',
            s_style))

        doc.build(elems); buf.seek(0)
        fname = f'Attendance_{subj.code}_{calendar.month_name[month]}_{year}.pdf'
        return send_file(buf, as_attachment=True, download_name=fname, mimetype='application/pdf')

    # ── Admin — Users ────────────────────────────────────────────────────────
    @app.route('/admin/users')
    @login_required
    @admin_required
    def admin_users():
        return render_template('admin/users.html',
            users=User.query.order_by(User.role, User.name).all(),
            departments=Department.query.all())

    @app.route('/admin/users/add', methods=['POST'])
    @login_required
    @admin_required
    def admin_add_user():
        u = User(
            username    = request.form['username'],
            email       = request.form['email'],
            name        = request.form['name'],
            role        = request.form['role'],
            department_id = request.form.get('department_id') or None)
        u.set_password(request.form['password'])
        db.session.add(u); db.session.flush()
        if u.role == 'student':
            db.session.add(Student(
                usn=request.form['usn'], name=u.name,
                department_id=u.department_id,
                semester=int(request.form['semester']),
                section=request.form.get('section', 'A'),
                user_id=u.id, email=u.email))
        db.session.commit()
        flash(f'User {u.name} added.', 'success')
        return redirect(url_for('admin_users'))

    @app.route('/admin/users/toggle/<int:user_id>', methods=['POST'])
    @login_required
    @admin_required
    def admin_toggle_user(user_id):
        u = db.session.get(User, user_id) or abort(404)
        u.is_active_user = not u.is_active_user
        db.session.commit()
        flash(f'User {"activated" if u.is_active_user else "deactivated"}.', 'success')
        return redirect(url_for('admin_users'))

    # ── Admin — Subjects ─────────────────────────────────────────────────────
    @app.route('/admin/subjects')
    @login_required
    @admin_required
    def admin_subjects():
        dept_id  = request.args.get('dept_id', type=int)
        subjects = (Subject.query.filter_by(department_id=dept_id).all()
                    if dept_id else Subject.query.all())
        return render_template('admin/subjects.html',
            subjects=subjects,
            departments=Department.query.all(),
            faculty_users=User.query.filter(User.role.in_(['faculty','hod'])).all(),
            dept_id=dept_id)

    @app.route('/admin/subjects/add', methods=['POST'])
    @login_required
    @admin_required
    def admin_add_subject():
        s = Subject(
            name=request.form['name'], code=request.form['code'],
            department_id=int(request.form['department_id']),
            semester=int(request.form['semester']),
            credits=int(request.form.get('credits', 4)),
            faculty_id=request.form.get('faculty_id') or None,
            section=request.form.get('section', 'A'))
        db.session.add(s); db.session.commit()
        flash(f'Subject "{s.name}" added.', 'success')
        return redirect(url_for('admin_subjects'))

    # ── JSON APIs ────────────────────────────────────────────────────────────
    @app.route('/api/subjects')
    @login_required
    def api_subjects():
        dept_id  = request.args.get('dept_id',  type=int)
        semester = request.args.get('semester', type=int)
        q = Subject.query
        if dept_id:  q = q.filter_by(department_id=dept_id)
        if semester: q = q.filter_by(semester=semester)
        return jsonify([{'id': s.id, 'name': s.name, 'code': s.code, 'section': s.section}
                        for s in q.all()])

    # ── Error Handlers ───────────────────────────────────────────────────────
    @app.errorhandler(403)
    def forbidden(e):
        return render_template('errors/403.html'), 403

    @app.errorhandler(404)
    def not_found(e):
        return render_template('errors/404.html'), 404

    with app.app_context():
        db.create_all()

    return app


app = create_app()

if __name__ == '__main__':
    app.run(host=Config.HOST, port=Config.PORT, debug=True)
